"""Genera i due fogli di calcolo demo per Giorgio Scala.

  Cash_Flow_Previsionale_2027.xlsx
  Calcolatore_IMU_2027.xlsx

Formule limitate a funzioni supportate da Odoo Spreadsheet (SUMIFS, SUMPRODUCT,
VLOOKUP, PMT, FV, EDATE, ...), così i file si possono caricare in Documenti e
aprire come foglio Odoo.

Uso: python3 build_spreadsheets.py <cartella_output>

Dopo la generazione, ri-salvare i file con LibreOffice (calcola i valori e
scrive percorsi relativi, necessari all'import xlsx di Odoo):
  soffice --headless --calc --convert-to xlsx --outdir <dir> <file>.xlsx
"""
import os
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from demo_data import (ANNO, CATEGORIE, COMUNI, IMMOBILI, MUTUI, SOCIETA,  # noqa: E402
                       STRAORDINARI)

PURPLE = "714B67"  # viola Odoo
F_TITLE = Font(bold=True, size=16, color=PURPLE)
F_SUB = Font(italic=True, size=10, color="666666")
F_HEAD = Font(bold=True, color="FFFFFF")
F_BOLD = Font(bold=True)
FILL_HEAD = PatternFill("solid", fgColor=PURPLE)
FILL_INPUT = PatternFill("solid", fgColor="FFF8DC")
FILL_SECTION = PatternFill("solid", fgColor="EDE7EC")
FILL_TOTAL = PatternFill("solid", fgColor="D9D2D8")
THIN = Side(style="thin", color="BBBBBB")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
EUR = '#,##0[$€]'
EUR2 = '#,##0.00[$€]'
PCT = "0.00%"
PCT1 = "0.0%"
RED_FONT = Font(color="C00000", bold=True)

N_IMM = len(IMMOBILI)
MONTHS_IT = ["Gen", "Feb", "Mar", "Apr", "Mag", "Giu", "Lug", "Ago", "Set", "Ott", "Nov", "Dic"]


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F_SUB


def header(ws, row, labels, start_col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=lab)
        c.font = F_HEAD
        c.fill = FILL_HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = 32


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[col(i)].width = w


def cell(ws, ref, value, fmt=None, inp=False, bold=False):
    c = ws[ref]
    c.value = value
    if fmt:
        c.number_format = fmt
    if inp:
        c.fill = FILL_INPUT
    if bold:
        c.font = F_BOLD
    c.border = BOX
    return c


def instructions(ws, heading, lines):
    title(ws, heading)
    ws.column_dimensions["A"].width = 120
    r = 3
    for line in lines:
        c = ws.cell(row=r, column=1, value=line)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if line and line.isupper():
            c.font = Font(bold=True, color=PURPLE)
        r += 1
    ws.cell(row=r + 1, column=1, value="Celle gialle = dati da inserire/modificare. Celle bianche = formule.").fill = FILL_INPUT


# ---------------------------------------------------------------------------
# CASH FLOW
# ---------------------------------------------------------------------------

def build_cash_flow(path):
    wb = Workbook()
    ist = wb.active
    ist.title = "Istruzioni"
    instructions(ist, f"Cash Flow previsionale {ANNO} – Immobiliare Uno & Immobiliare Due", [
        "OBIETTIVO",
        "Previsione dei flussi di cassa mensili delle due società (e consolidata) e proiezione a 5 anni, "
        "con dati per singolo immobile: canoni, sfitti, morosità, IMU, imposta di registro, mutui, "
        "manutenzioni ordinarie/straordinarie, imprevisti, compravendite.",
        "",
        "COME SI USA",
        "1. Parametri: scegli la società nel filtro (Tutte / IU / ID), anno, saldi iniziali, % imprevisti e morosità.",
        "2. Immobili: anagrafica del patrimonio con canone, rendita catastale, valore di mercato, costi ricorrenti. "
        "Il foglio calcola IMU, registro, rendimento lordo/netto e incidenza costi per immobile.",
        "3. Occupazione: 1 = locato / 0 = sfitto per mese. Alimenta canoni e IMU (beni merce: IMU solo sui mesi locati).",
        "4. Mutui: debito residuo, TAN e rate residue. Calcola rata, interessi, quota capitale e simulazione surroga.",
        "5. Straordinari: movimenti spot (ritinteggiature, guasti, rateizzazioni arretrati, acquisti/vendite immobili).",
        "6. Cash Flow: risultato mensile; Proiezione 5 anni: medio-lungo termine; KPI Patrimonio: sintesi per società.",
        "",
        "COLLEGAMENTO CON ODOO",
        "In Odoo il file si apre da Documenti come foglio di calcolo. Le righe consuntive possono essere sostituite "
        "con formule ODOO.BALANCE / ODOO.PIVOT (es. incassi reali dal registro Banca, fatture canoni da Abbonamenti, "
        "rate mutui dal modulo Prestiti) per confrontare previsto vs consuntivo.",
        "",
        "NOTE FISCALI",
        "Locazioni in esenzione IVA art. 10 DPR 633/72. Imposta di registro annuale: 1% strumentali, 2% abitativi "
        "(da versare entro 30 gg dalla ricorrenza). IMU: acconto 16/06 e saldo 16/12; per i beni merce "
        "(fabbricati dell'impresa costruttrice) esenzione nei mesi non locati. Valori demo da verificare con il commercialista.",
    ])

    # ---- Parametri ----
    p = wb.create_sheet("Parametri")
    title(p, "Parametri generali")
    widths(p, [34, 26, 16, 18, 16, 18, 10, 3, 12, 14, 3, 14, 14, 14])
    cell(p, "A3", "Filtro società (Tutte / IU / ID)", bold=True)
    cell(p, "B3", "Tutte", inp=True)
    dv = DataValidation(type="list", formula1='"Tutte,IU,ID"', allow_blank=False)
    p.add_data_validation(dv)
    dv.add("B3")
    cell(p, "A4", "Anno di previsione", bold=True)
    cell(p, "B4", ANNO, inp=True)

    header(p, 6, ["Codice", "Ragione sociale", "Saldo cassa 01/01", "IRES+IRAP annua stimata",
                  "Assicurazioni annue", "Commercialista / consulenze mensili", "Incluso"])
    for i, (code, name, saldo, ires, ass, comm) in enumerate(SOCIETA):
        r = 7 + i
        cell(p, f"A{r}", code, inp=True)
        cell(p, f"B{r}", name, inp=True)
        cell(p, f"C{r}", saldo, EUR, inp=True)
        cell(p, f"D{r}", ires, EUR, inp=True)
        cell(p, f"E{r}", ass, EUR, inp=True)
        cell(p, f"F{r}", comm, EUR, inp=True)
        cell(p, f"G{r}", f'=IF(OR($B$3="Tutte",A{r}=$B$3),1,0)')
    soc_last = 6 + len(SOCIETA)

    params = [
        ("% imprevisti (accantonamento su canoni)", 0.03, PCT1),
        ("% morosità attesa sui canoni", 0.02, PCT1),
        ("Mese pagamento assicurazioni (1-12)", 1, None),
        ("% acconto imposte dirette a giugno", 0.40, PCT1),
        ("% acconto imposte dirette a novembre", 0.60, PCT1),
        ("Adeguamento ISTAT canoni annuo (proiezione)", 0.02, PCT1),
        ("Inflazione costi annua (proiezione)", 0.02, PCT1),
        ("Riserva manut. straordinaria (% valore patrimonio / anno)", 0.005, PCT),
    ]
    for i, (lab, val, fmt) in enumerate(params):
        r = 11 + i
        cell(p, f"A{r}", lab, bold=True)
        cell(p, f"B{r}", val, fmt, inp=True)
    # B11 imprevisti, B12 morosità, B13 mese ass., B14 giugno, B15 novembre, B16 ISTAT, B17 inflazione, B18 riserva

    header(p, 6, ["Categoria", "Moltipl."], start_col=9)
    for i, (cat, molt, _g, _ct) in enumerate(CATEGORIE):
        r = 7 + i
        cell(p, f"I{r}", cat, inp=True)
        cell(p, f"J{r}", molt, inp=True)
    cat_rng = f"Parametri!$I$7:$J${6 + len(CATEGORIE)}"

    header(p, 6, ["Comune", "Aliquota IMU anno prec. ‰", "Aliquota IMU anno corr. ‰"], start_col=12)
    for i, (com, _cc, a_prev, a_corr) in enumerate(COMUNI):
        r = 7 + i
        cell(p, f"L{r}", com, inp=True)
        cell(p, f"M{r}", a_prev, "0.00", inp=True)
        cell(p, f"N{r}", a_corr, "0.00", inp=True)
    com_rng = f"Parametri!$L$7:$N${6 + len(COMUNI)}"
    p.freeze_panes = "A6"

    # ---- Immobili ----
    im = wb.create_sheet("Immobili")
    title(im, "Anagrafica immobili e redditività", "Dati demo – IMU calcolata su rendita rivalutata 5% × moltiplicatore × aliquota comunale")
    heads = ["Codice", "Società", "Descrizione", "Comune", "Tipo", "Cat.", "Rendita catastale",
             "Valore di mercato", "Canone mensile", "Inquilino", "Mese rinnovo registro", "% registro",
             "Condominio mensile", "Manut. ordinaria annua", "Bene merce", "Mesi locati",
             "Mesi locati 1° sem.", "Canone annuo effettivo", "Base imponibile IMU", "IMU acconto (giugno)",
             "IMU annua", "Registro annuo", "Manut. straordinaria anno", "Costi totali anno",
             "Rend. lordo potenziale", "Rend. lordo effettivo", "Rend. netto", "Incidenza costi su canoni",
             "Incluso"]
    header(im, 4, heads)
    widths(im, [8, 7, 32, 11, 12, 6, 11, 13, 11, 28, 9, 8, 11, 11, 8, 7, 8, 13, 13, 12, 11, 11, 12, 12,
                10, 10, 10, 11, 7])
    first, last = 5, 4 + N_IMM
    for i, rec in enumerate(IMMOBILI):
        (code, soc, desc, com, tipo, cat, rend, val, can, inq, anniv, reg, cond, man, merce, _occ) = rec
        r = first + i
        for c_, v, fmt in [("A", code, None), ("B", soc, None), ("C", desc, None), ("D", com, None),
                           ("E", tipo, None), ("F", cat, None), ("G", rend, EUR2), ("H", val, EUR),
                           ("I", can, EUR), ("J", inq, None), ("K", anniv, None), ("L", reg, PCT1),
                           ("M", cond, EUR), ("N", man, EUR), ("O", merce, None)]:
            cell(im, f"{c_}{r}", v, fmt, inp=True)
        cell(im, f"P{r}", f"=SUM(Occupazione!C{r}:N{r})")
        cell(im, f"Q{r}", f"=SUM(Occupazione!C{r}:H{r})")
        cell(im, f"R{r}", f"=I{r}*P{r}", EUR)
        cell(im, f"S{r}", f"=G{r}*1.05*VLOOKUP(F{r},{cat_rng},2,0)", EUR)
        cell(im, f"T{r}", f'=ROUND(S{r}*VLOOKUP(D{r},{com_rng},2,0)/1000*IF(O{r}="Sì",Q{r},6)/12,2)', EUR2)
        cell(im, f"U{r}", f'=ROUND(S{r}*VLOOKUP(D{r},{com_rng},3,0)/1000*IF(O{r}="Sì",P{r},12)/12,2)', EUR2)
        cell(im, f"V{r}", f"=IF(P{r}>0,I{r}*12*L{r},0)", EUR)
        cell(im, f"W{r}", f'=-SUMIFS(Straordinari!$F$5:$F$64,Straordinari!$C$5:$C$64,A{r},'
                          f'Straordinari!$D$5:$D$64,"Manutenzione straordinaria")', EUR)
        cell(im, f"X{r}", f"=U{r}+V{r}+M{r}*12+N{r}+W{r}", EUR)
        cell(im, f"Y{r}", f"=IFERROR(I{r}*12/H{r},0)", PCT)
        cell(im, f"Z{r}", f"=IFERROR(R{r}/H{r},0)", PCT)
        cell(im, f"AA{r}", f"=IFERROR((R{r}-X{r})/H{r},0)", PCT)
        cell(im, f"AB{r}", f"=IFERROR(X{r}/R{r},0)", PCT1)
        cell(im, f"AC{r}", f'=IF(OR(Parametri!$B$3="Tutte",B{r}=Parametri!$B$3),1,0)')
    tr = last + 1
    cell(im, f"A{tr}", "TOTALE (filtro)", bold=True)
    for c_ in ["H", "R", "T", "U", "V", "W", "X"]:
        cell(im, f"{c_}{tr}", f"=SUMPRODUCT($AC${first}:$AC${last},{c_}{first}:{c_}{last})", EUR, bold=True)
    cell(im, f"I{tr}", f"=SUMPRODUCT($AC${first}:$AC${last},I{first}:I{last})", EUR, bold=True)
    cell(im, f"Y{tr}", f"=IFERROR(I{tr}*12/H{tr},0)", PCT, bold=True)
    cell(im, f"Z{tr}", f"=IFERROR(R{tr}/H{tr},0)", PCT, bold=True)
    cell(im, f"AA{tr}", f"=IFERROR((R{tr}-X{tr})/H{tr},0)", PCT, bold=True)
    cell(im, f"AB{tr}", f"=IFERROR(X{tr}/R{tr},0)", PCT1, bold=True)
    for c_ in range(1, 30):
        im.cell(row=tr, column=c_).fill = FILL_TOTAL
    im.freeze_panes = "D5"
    dv2 = DataValidation(type="list", formula1='"Sì,No"')
    im.add_data_validation(dv2)
    dv2.add(f"O{first}:O{last + 20}")

    # ---- Occupazione ----
    oc = wb.create_sheet("Occupazione")
    title(oc, f"Piano di occupazione {ANNO}", "1 = locato, 0 = sfitto")
    header(oc, 4, ["Codice", "Inquilino"] + MONTHS_IT + ["Mesi locati"])
    widths(oc, [9, 30] + [6] * 12 + [10])
    for i, rec in enumerate(IMMOBILI):
        r = first + i
        cell(oc, f"A{r}", f"=Immobili!A{r}")
        cell(oc, f"B{r}", f"=Immobili!J{r}")
        for m, v in enumerate(rec[15]):
            cell(oc, f"{col(3 + m)}{r}", v, inp=True).alignment = Alignment(horizontal="center")
        cell(oc, f"O{r}", f"=SUM(C{r}:N{r})", bold=True)
    oc.conditional_formatting.add(f"C{first}:N{last}",
                                  CellIsRule(operator="equal", formula=["0"], fill=PatternFill("solid", fgColor="F4CCCC")))
    oc.freeze_panes = "C5"

    # ---- Mutui ----
    mu = wb.create_sheet("Mutui")
    title(mu, "Mutui e finanziamenti", "Rata francese; simulazione surroga/rinegoziazione sul debito residuo")
    header(mu, 4, ["Società", "Descrizione", "Banca", "Immobili a garanzia", "Debito residuo 01/01",
                   "TAN attuale", "Rate residue", "Rata mensile", "Debito residuo 31/12",
                   "Quota capitale anno", "Interessi anno", "TAN ipotesi surroga", "Rata con surroga",
                   "Risparmio annuo surroga", "Incluso", "Rate anno 1", "Rate anno 2", "Rate anno 3",
                   "Rate anno 4", "Rate anno 5"])
    widths(mu, [8, 30, 18, 20, 14, 9, 9, 12, 14, 13, 12, 10, 12, 13, 8, 12, 12, 12, 12, 12])
    m_first = 5
    m_last = 4 + len(MUTUI)
    for i, (soc, desc, banca, gar, res, tan, n, tan_s) in enumerate(MUTUI):
        r = m_first + i
        for c_, v, fmt in [("A", soc, None), ("B", desc, None), ("C", banca, None), ("D", gar, None),
                           ("E", res, EUR), ("F", tan, PCT), ("G", n, None), ("L", tan_s, PCT)]:
            cell(mu, f"{c_}{r}", v, fmt, inp=True)
        cell(mu, f"H{r}", f"=PMT(F{r}/12,G{r},-E{r})", EUR2)
        cell(mu, f"I{r}", f"=-FV(F{r}/12,MIN(12,G{r}),-H{r},E{r})", EUR)
        cell(mu, f"J{r}", f"=E{r}-I{r}", EUR)
        cell(mu, f"K{r}", f"=H{r}*MIN(12,G{r})-J{r}", EUR)
        cell(mu, f"M{r}", f"=PMT(L{r}/12,G{r},-E{r})", EUR2)
        cell(mu, f"N{r}", f"=(H{r}-M{r})*12", EUR)
        cell(mu, f"O{r}", f'=IF(OR(Parametri!$B$3="Tutte",A{r}=Parametri!$B$3),1,0)')
        for k in range(5):
            cell(mu, f"{col(16 + k)}{r}", f"=H{r}*MAX(0,MIN(12,G{r}-12*{k}))", EUR)
    tr_m = m_last + 1
    cell(mu, f"A{tr_m}", "TOTALE (filtro)", bold=True)
    for c_ in ["E", "H", "I", "J", "K", "M", "N", "P", "Q", "R", "S", "T"]:
        cell(mu, f"{c_}{tr_m}", f"=SUMPRODUCT($O${m_first}:$O${m_last},{c_}{m_first}:{c_}{m_last})", EUR, bold=True)
    for c_ in range(1, 21):
        mu.cell(row=tr_m, column=c_).fill = FILL_TOTAL
    mu.freeze_panes = "C5"

    # ---- Straordinari ----
    st = wb.create_sheet("Straordinari")
    title(st, "Movimenti straordinari previsti", "Importi + entrata / − uscita. Categorie: Manutenzione straordinaria, "
          "Rateizzazione arretrati, Compravendita, Commerciale, Altro")
    header(st, 4, ["Data", "Società", "Immobile", "Categoria", "Descrizione", "Importo", "Incluso"])
    widths(st, [12, 9, 10, 26, 44, 13, 8])
    dv3 = DataValidation(type="list", formula1='"Manutenzione straordinaria,Rateizzazione arretrati,Compravendita,Commerciale,Altro"')
    st.add_data_validation(dv3)
    dv3.add("D5:D64")
    for i in range(60):
        r = 5 + i
        if i < len(STRAORDINARI):
            d, soc, imm, cat, desc, imp = STRAORDINARI[i]
            y, mo, dd = map(int, d.split("-"))
            cell(st, f"A{r}", date(y, mo, dd), "dd/mm/yyyy", inp=True)
            cell(st, f"B{r}", soc, inp=True)
            cell(st, f"C{r}", imm, inp=True)
            cell(st, f"D{r}", cat, inp=True)
            cell(st, f"E{r}", desc, inp=True)
            cell(st, f"F{r}", imp, EUR, inp=True)
        else:
            for c_ in "ABCDEF":
                cell(st, f"{c_}{r}", None, "dd/mm/yyyy" if c_ == "A" else (EUR if c_ == "F" else None), inp=True)
        cell(st, f"G{r}", f'=IF(B{r}="",0,IF(OR(Parametri!$B$3="Tutte",B{r}=Parametri!$B$3),1,0))')
    st.freeze_panes = "A5"

    # ---- Cash Flow ----
    cf = wb.create_sheet("Cash Flow", 1)
    title(cf, f"Cash Flow previsionale {ANNO}")
    cf["A2"] = '="Società: "&Parametri!B3&"   (cambia il filtro nel foglio Parametri)"'
    cf["A2"].font = F_SUB
    widths(cf, [44] + [12] * 12 + [14])
    header(cf, 4, ["Voce"] + [None] * 12 + ["Totale anno"])
    for m in range(12):
        c = cf.cell(row=4, column=2 + m)
        c.value = f"=DATE(Parametri!$B$4,{m + 1},1)"
        c.number_format = "mmm yy"
    imm = lambda c_: f"Immobili!${c_}${first}:${c_}${last}"  # noqa: E731
    inc_imm = imm("AC")
    stf = "Straordinari!$F$5:$F$64"
    sta = "Straordinari!$A$5:$A$64"
    std = "Straordinari!$D$5:$D$64"
    stg = "Straordinari!$G$5:$G$64"

    def st_sum(cat, k, sign=None):
        h = f"{col(k)}$4"
        s = f'SUMIFS({stf},{stg},1,{std},"{cat}",{sta},">="&{h},{sta},"<"&EDATE({h},1)'
        if sign:
            s += f',{stf},"{sign}0"'
        return s + ")"

    rows = [
        ("ENTRATE", None),
        ("Canoni di locazione (esenti IVA art.10)",
         lambda k, oc_col: f"=SUMPRODUCT({inc_imm},{imm('I')},Occupazione!{oc_col}${first}:{oc_col}${last})"),
        ("Morosità attesa", lambda k, oc_col: f"=-{col(k)}6*Parametri!$B$12"),
        ("Rateizzazioni arretrati incassate", lambda k, oc_col: "=" + st_sum("Rateizzazione arretrati", k)),
        ("Vendite immobili", lambda k, oc_col: "=" + st_sum("Compravendita", k, ">")),
        ("Totale entrate", lambda k, oc_col: f"=SUM({col(k)}6:{col(k)}9)"),
        (None, None),
        ("USCITE", None),
        ("Rate mutui e finanziamenti",
         lambda k, oc_col: f'=-SUMIFS(Mutui!$H${m_first}:$H${m_last},Mutui!$O${m_first}:$O${m_last},1,'
                           f'Mutui!$G${m_first}:$G${m_last},">="&{k - 1})'),
        ("IMU (acconto 16/06 – saldo 16/12)",
         lambda k, oc_col: f"=IF(MONTH({col(k)}$4)=6,-SUMIFS({imm('T')},{inc_imm},1),"
                           f"IF(MONTH({col(k)}$4)=12,-(SUMIFS({imm('U')},{inc_imm},1)-SUMIFS({imm('T')},{inc_imm},1)),0))"),
        ("Imposta di registro annuale",
         lambda k, oc_col: f"=-SUMIFS({imm('V')},{inc_imm},1,{imm('K')},MONTH({col(k)}$4))"),
        ("Spese condominiali (quota proprietà)", lambda k, oc_col: f"=-SUMIFS({imm('M')},{inc_imm},1)"),
        ("Manutenzione ordinaria programmata", lambda k, oc_col: f"=-SUMIFS({imm('N')},{inc_imm},1)/12"),
        ("Manutenzioni straordinarie / ritinteggiature",
         lambda k, oc_col: "=" + st_sum("Manutenzione straordinaria", k)),
        ("Imprevisti (accantonamento)", lambda k, oc_col: f"=-{col(k)}6*Parametri!$B$11"),
        ("Assicurazioni",
         lambda k, oc_col: f"=IF(MONTH({col(k)}$4)=Parametri!$B$13,-SUMIFS(Parametri!$E$7:$E${soc_last},"
                           f"Parametri!$G$7:$G${soc_last},1),0)"),
        ("Commercialista e consulenze",
         lambda k, oc_col: f"=-SUMIFS(Parametri!$F$7:$F${soc_last},Parametri!$G$7:$G${soc_last},1)"),
        ("Imposte dirette (IRES/IRAP)",
         lambda k, oc_col: f"=-SUMIFS(Parametri!$D$7:$D${soc_last},Parametri!$G$7:$G${soc_last},1)*"
                           f"IF(MONTH({col(k)}$4)=6,Parametri!$B$14,IF(MONTH({col(k)}$4)=11,Parametri!$B$15,0))"),
        ("Spese commerciali (annunci, agenzie)", lambda k, oc_col: "=" + st_sum("Commerciale", k)),
        ("Investimenti / acquisti immobili", lambda k, oc_col: "=" + st_sum("Compravendita", k, "<")),
        ("Totale uscite", lambda k, oc_col: f"=SUM({col(k)}13:{col(k)}24)"),
        (None, None),
        ("FLUSSO DI CASSA NETTO", lambda k, oc_col: f"={col(k)}10+{col(k)}25"),
        ("Saldo cassa iniziale",
         lambda k, oc_col: (f"=SUMIFS(Parametri!$C$7:$C${soc_last},Parametri!$G$7:$G${soc_last},1)"
                            if k == 2 else f"={col(k - 1)}29")),
        ("SALDO CASSA FINALE", lambda k, oc_col: f"={col(k)}28+{col(k)}27"),
    ]
    for i, (lab, fn) in enumerate(rows):
        r = 5 + i
        if lab is None:
            continue
        a = cf.cell(row=r, column=1, value=lab)
        a.border = BOX
        if fn is None:
            for k in range(1, 15):
                cf.cell(row=r, column=k).fill = FILL_SECTION
            a.font = Font(bold=True, color=PURPLE)
            continue
        is_total = lab.startswith("Totale") or lab.isupper()
        for m in range(12):
            k = 2 + m
            c = cf.cell(row=r, column=k, value=fn(k, col(3 + m)))
            c.number_format = EUR
            c.border = BOX
            if is_total:
                c.font = F_BOLD
                c.fill = FILL_TOTAL
        tot = cf.cell(row=r, column=14)
        if r == 28:
            tot.value = "=B28"
        elif r == 29:
            tot.value = "=M29"
        else:
            tot.value = f"=SUM(B{r}:M{r})"
        tot.number_format = EUR
        tot.font = F_BOLD
        tot.border = BOX
        tot.fill = FILL_TOTAL
        if is_total:
            a.font = F_BOLD
            a.fill = FILL_TOTAL
    cf.conditional_formatting.add("B27:N29", CellIsRule(operator="lessThan", formula=["0"], font=RED_FONT,
                                                         fill=PatternFill("solid", fgColor="F4CCCC")))
    cf.freeze_panes = "B5"

    ch = LineChart()
    ch.title = "Saldo di cassa mensile"
    ch.height, ch.width = 8, 22
    ch.add_data(Reference(cf, min_col=1, max_col=13, min_row=29, max_row=29), from_rows=True, titles_from_data=True)
    ch.set_categories(Reference(cf, min_col=2, max_col=13, min_row=4, max_row=4))
    ch.y_axis.number_format = '#,##0'
    cf.add_chart(ch, "B32")
    bc = BarChart()
    bc.title = "Entrate vs uscite"
    bc.height, bc.width = 8, 22
    bc.add_data(Reference(cf, min_col=1, max_col=13, min_row=10, max_row=10), from_rows=True, titles_from_data=True)
    bc.add_data(Reference(cf, min_col=1, max_col=13, min_row=25, max_row=25), from_rows=True, titles_from_data=True)
    bc.set_categories(Reference(cf, min_col=2, max_col=13, min_row=4, max_row=4))
    cf.add_chart(bc, "B49")

    # ---- Proiezione 5 anni ----
    pj = wb.create_sheet("Proiezione 5 anni", 2)
    title(pj, "Proiezione di medio-lungo termine (5 anni)",
          "Anno 1 = Cash Flow mensile; anni successivi con ISTAT, inflazione costi, rate mutui residue e riserva manutenzioni")
    widths(pj, [44, 14, 14, 14, 14, 14, 15])
    header(pj, 4, ["Voce"] + [f"=Parametri!$B$4+{k}" for k in range(5)] + ["Totale 5 anni"])
    for k in range(5):
        pj.cell(row=4, column=2 + k).number_format = "0"
    prj = [
        ("Canoni netti (dopo morosità)", lambda k: "='Cash Flow'!N6+'Cash Flow'!N7" if k == 0
         else f"=B5*(1+Parametri!$B$16)^{k}"),
        ("Entrate straordinarie (rateizzazioni, vendite)", lambda k: "='Cash Flow'!N8+'Cash Flow'!N9" if k == 0 else "=0"),
        ("Rate mutui", lambda k: f"=-SUMIFS(Mutui!${col(16 + k)}${m_first}:${col(16 + k)}${m_last},"
                                 f"Mutui!$O${m_first}:$O${m_last},1)"),
        ("IMU", lambda k: "='Cash Flow'!N14" if k == 0 else "=$B$8"),
        ("Costi operativi ricorrenti", lambda k: "='Cash Flow'!N15+'Cash Flow'!N16+'Cash Flow'!N17+'Cash Flow'!N19+"
                                                 "'Cash Flow'!N20+'Cash Flow'!N21+'Cash Flow'!N22" if k == 0
         else f"=$B$9*(1+Parametri!$B$17)^{k}"),
        ("Manutenzioni straordinarie", lambda k: "='Cash Flow'!N18" if k == 0
         else f"=-Immobili!$H${tr}*Parametri!$B$18*(1+Parametri!$B$17)^{k}"),
        ("Investimenti e spese commerciali", lambda k: "='Cash Flow'!N23+'Cash Flow'!N24" if k == 0 else "=0"),
        ("FLUSSO NETTO", lambda k: f"=SUM({col(2 + k)}5:{col(2 + k)}11)"),
        ("Saldo cassa iniziale", lambda k: "='Cash Flow'!B28" if k == 0 else f"={col(1 + k)}14"),
        ("SALDO CASSA FINALE", lambda k: f"={col(2 + k)}13+{col(2 + k)}12"),
        ("Debito residuo mutui fine anno (stima)", lambda k: f"=Mutui!$I${tr_m}" if k == 0
         else f"=MAX(0,{col(1 + k)}15-(-{col(2 + k)}7)*0.7)"),
    ]
    for i, (lab, fn) in enumerate(prj):
        r = 5 + i
        a = cell(pj, f"A{r}", lab)
        bold = lab.isupper()
        if bold:
            a.font = F_BOLD
            a.fill = FILL_TOTAL
        for k in range(5):
            c = cell(pj, f"{col(2 + k)}{r}", fn(k), EUR, bold=bold)
            if bold:
                c.fill = FILL_TOTAL
        if r in (13, 14, 15):
            cell(pj, f"G{r}", "=F14" if r == 14 else ("=B13" if r == 13 else "=F15"), EUR, bold=True)
        else:
            cell(pj, f"G{r}", f"=SUM(B{r}:F{r})", EUR, bold=True)
    pj["A17"] = "Nota: il debito residuo dagli anni successivi è stimato (quota capitale ≈ 70% delle rate)."
    pj["A17"].font = F_SUB
    pj.conditional_formatting.add("B12:G14", CellIsRule(operator="lessThan", formula=["0"], font=RED_FONT))
    lc = LineChart()
    lc.title = "Saldo di cassa a 5 anni"
    lc.height, lc.width = 8, 18
    lc.add_data(Reference(pj, min_col=1, max_col=6, min_row=14, max_row=14), from_rows=True, titles_from_data=True)
    lc.set_categories(Reference(pj, min_col=2, max_col=6, min_row=4, max_row=4))
    pj.add_chart(lc, "A19")

    # ---- KPI Patrimonio ----
    kp = wb.create_sheet("KPI Patrimonio", 3)
    title(kp, "KPI del patrimonio immobiliare", "Per società e totale (indipendente dal filtro)")
    widths(kp, [40, 18, 18, 18])
    header(kp, 4, ["Indicatore"] + [s[0] for s in SOCIETA] + ["Totale"])
    socs = [s[0] for s in SOCIETA]

    def by_soc(expr_col, sc):
        return f'SUMIFS({imm(expr_col)},{imm("B")},"{sc}")'

    S = "sum"
    kpis = [
        ("Numero immobili", lambda sc: f'=COUNTIFS({imm("B")},"{sc}")', S, "0"),
        ("Valore di mercato patrimonio", lambda sc: "=" + by_soc("H", sc), S, EUR),
        ("Canone annuo potenziale", lambda sc: "=" + by_soc("I", sc) + "*12", S, EUR),
        ("Canone annuo effettivo (dopo sfitti)", lambda sc: "=" + by_soc("R", sc), S, EUR),
        ("Tasso di occupazione", lambda sc: f'=IFERROR({by_soc("P", sc)}/(COUNTIFS({imm("B")},"{sc}")*12),0)',
         f'=IFERROR(SUM({imm("P")})/({{c}}5*12),0)', PCT1),
        ("Mancato incasso per sfitto", lambda sc: "={c}7-{c}8", S, EUR),
        ("IMU annua", lambda sc: "=" + by_soc("U", sc), S, EUR),
        ("Imposta di registro", lambda sc: "=" + by_soc("V", sc), S, EUR),
        ("Manutenzioni (ord. + straord.)", lambda sc: f'={by_soc("N", sc)}+{by_soc("W", sc)}', S, EUR),
        ("Costi totali immobili", lambda sc: "=" + by_soc("X", sc), S, EUR),
        ("Rendimento lordo potenziale", lambda sc: "=IFERROR({c}7/{c}6,0)", None, PCT),
        ("Rendimento lordo effettivo", lambda sc: "=IFERROR({c}8/{c}6,0)", None, PCT),
        ("Rendimento netto (ante imposte dirette e oneri finanziari)",
         lambda sc: "=IFERROR(({c}8-{c}14)/{c}6,0)", None, PCT),
        ("Incidenza manutenzioni su canoni", lambda sc: "=IFERROR({c}13/{c}8,0)", None, PCT1),
        ("Debito residuo mutui 01/01",
         lambda sc: f'=SUMIFS(Mutui!$E${m_first}:$E${m_last},Mutui!$A${m_first}:$A${m_last},"{sc}")', S, EUR),
        ("Interessi passivi anno",
         lambda sc: f'=SUMIFS(Mutui!$K${m_first}:$K${m_last},Mutui!$A${m_first}:$A${m_last},"{sc}")', S, EUR),
        ("LTV (debito / valore)", lambda sc: "=IFERROR({c}19/{c}6,0)", None, PCT1),
        ("Rendimento netto dopo interessi", lambda sc: "=IFERROR(({c}8-{c}14-{c}20)/{c}6,0)", None, PCT),
    ]
    tc = col(2 + len(socs))
    for i, (lab, fn, total, fmt) in enumerate(kpis):
        r = 5 + i
        cell(kp, f"A{r}", lab, bold=True)
        for j, sc in enumerate(socs):
            c_ = col(2 + j)
            cell(kp, f"{c_}{r}", fn(sc).replace("{c}", c_), fmt)
        if total == S:
            tf = f"=SUM(B{r}:{col(1 + len(socs))}{r})"
        elif total is None:
            tf = fn("").replace("{c}", tc)
        else:
            tf = total.replace("{c}", tc)
        cell(kp, f"{tc}{r}", tf, fmt, bold=True)
    kp.freeze_panes = "B5"

    wb.save(path)


# ---------------------------------------------------------------------------
# IMU
# ---------------------------------------------------------------------------

def build_imu(path):
    wb = Workbook()
    ist = wb.active
    ist.title = "Istruzioni"
    instructions(ist, f"Calcolatore IMU {ANNO} – patrimonio Immobiliare Uno & Immobiliare Due", [
        "COME FUNZIONA",
        "Base imponibile = rendita catastale × 1,05 (rivalutazione) × moltiplicatore della categoria × (1 − riduzione).",
        "IMU = base imponibile × quota di possesso × aliquota comunale (‰) × mesi imponibili / 12.",
        "Mese imponibile: si conta il mese se il possesso/la locazione dura più della metà dei giorni del mese.",
        "",
        "BENI MERCE (fabbricati costruiti dall'impresa e destinati alla vendita)",
        "Esenti IMU finché non locati (L. 160/2019 c. 751). Se 'Bene merce' = Sì, i mesi imponibili sono solo i "
        "mesi di locazione: è la logica della dichiarazione IMU che 'cristallizza' il dovuto sui mesi effettivi.",
        "",
        "ACCONTO E SALDO",
        "Acconto 16 giugno: primo semestre con aliquote dell'anno precedente. Saldo 16 dicembre: imposta annua con "
        "aliquote deliberate per l'anno in corso meno l'acconto (conguaglio). Dichiarazione IMU: entro il 30 giugno "
        "dell'anno successivo.",
        "",
        "IMMOBILI DEL GRUPPO D",
        "Quota Stato 7,6‰ (codice tributo 3925); l'eccedenza deliberata dal Comune va con codice 3930. "
        "Gli altri fabbricati usano il codice 3918.",
        "",
        "FOGLI",
        "Parametri: moltiplicatori per categoria, aliquote per Comune (anno precedente e corrente). "
        "Calcolo IMU: un rigo per immobile. Riepilogo F24: importi per società, Comune e codice tributo. "
        "Budget vs Dichiarazione: IMU piena (budget) vs dovuta (dichiarazione) e risparmio da sfitto. "
        "Calcolatore rapido: simulazione su un singolo immobile, come i calcolatori online.",
        "",
        "ATTENZIONE: aliquote demo. Verificare sempre la delibera del Comune sul portale del MEF (finanze.gov.it) "
        "e validare il calcolo con il commercialista.",
    ])

    p = wb.create_sheet("Parametri")
    title(p, "Parametri IMU")
    widths(p, [30, 12, 3, 10, 12, 9, 14, 3, 14, 12, 14, 14])
    cell(p, "A3", "Anno d'imposta", bold=True)
    cell(p, "B3", ANNO, inp=True)
    cell(p, "A4", "Quota Stato gruppo D (‰)", bold=True)
    cell(p, "B4", 7.6, "0.00", inp=True)
    cell(p, "A5", "Rivalutazione rendita", bold=True)
    cell(p, "B5", 0.05, PCT1, inp=True)
    header(p, 7, ["Categoria", "Moltipl.", "Gruppo", "Cod. tributo Comune"], start_col=4)
    for i, (cat, molt, grp, ct) in enumerate(CATEGORIE):
        r = 8 + i
        cell(p, f"D{r}", cat, inp=True)
        cell(p, f"E{r}", molt, inp=True)
        cell(p, f"F{r}", grp, inp=True)
        cell(p, f"G{r}", ct, inp=True)
    cat_rng = f"Parametri!$D$8:$G${7 + len(CATEGORIE)}"
    header(p, 7, ["Comune", "Cod. catastale", "Aliquota anno prec. ‰", "Aliquota anno corr. ‰"], start_col=9)
    for i, (com, cc, a_prev, a_corr) in enumerate(COMUNI):
        r = 8 + i
        cell(p, f"I{r}", com, inp=True)
        cell(p, f"J{r}", cc, inp=True)
        cell(p, f"K{r}", a_prev, "0.00", inp=True)
        cell(p, f"L{r}", a_corr, "0.00", inp=True)
    com_rng = f"Parametri!$I$8:$L${7 + len(COMUNI)}"

    ca = wb.create_sheet("Calcolo IMU", 1)
    title(ca, f"Calcolo IMU {ANNO} per immobile", "Le colonne gialle sono i parametri di ogni immobile")
    heads = ["Società", "Immobile", "Descrizione", "Comune", "Cat.", "Rendita catastale", "Quota possesso",
             "Bene merce", "Mesi possesso 1° sem", "Mesi possesso 2° sem", "Mesi locati 1° sem",
             "Mesi locati 2° sem", "Riduzione (inagibile/storico)", "Rendita rivalutata", "Moltipl.",
             "Base imponibile", "Aliquota acconto ‰", "Aliquota saldo ‰", "Mesi imponibili 1° sem",
             "Mesi imponibili 2° sem", "IMU piena (budget)", "Acconto 16/06", "IMU annua dovuta",
             "Saldo 16/12", "Risparmio vs budget", "Gruppo", "Cod. tributo Comune", "di cui quota Stato (3925)",
             "di cui quota Comune", "Acconto quota Stato", "Saldo quota Stato", "Cod. catastale"]
    header(ca, 4, heads)
    widths(ca, [7, 8, 30, 10, 6, 11, 8, 7, 8, 8, 8, 8, 9, 11, 7, 12, 8, 8, 8, 8, 12, 12, 12, 12, 12, 7, 9,
                12, 12, 11, 11, 9])
    first = 5
    n_rows = N_IMM + 10
    last = first + n_rows - 1
    for i in range(n_rows):
        r = first + i
        if i < N_IMM:
            (code, soc, desc, com, _t, cat, rend, *_rest) = IMMOBILI[i]
            occ = IMMOBILI[i][15]
            merce = IMMOBILI[i][14]
            vals = [soc, code, desc, com, cat, rend, 1, merce, 6, 6, sum(occ[:6]), sum(occ[6:]), 0]
        else:
            vals = [None] * 13
        fmts = [None, None, None, None, None, EUR2, "0%", None, "0", "0", "0", "0", "0%"]
        for j, (v, fmt) in enumerate(zip(vals, fmts)):
            cell(ca, f"{col(1 + j)}{r}", v, fmt, inp=True)
        g = f'IF($A{r}="",0,'
        cell(ca, f"N{r}", f"={g}F{r}*(1+Parametri!$B$5))", EUR2)
        cell(ca, f"O{r}", f"={g}VLOOKUP(E{r},{cat_rng},2,0))")
        cell(ca, f"P{r}", f"={g}N{r}*O{r}*(1-M{r}))", EUR2)
        cell(ca, f"Q{r}", f"={g}VLOOKUP(D{r},{com_rng},3,0))", "0.00")
        cell(ca, f"R{r}", f"={g}VLOOKUP(D{r},{com_rng},4,0))", "0.00")
        cell(ca, f"S{r}", f'={g}IF(H{r}="Sì",MIN(I{r},K{r}),I{r}))', "0")
        cell(ca, f"T{r}", f'={g}IF(H{r}="Sì",MIN(J{r},L{r}),J{r}))', "0")
        cell(ca, f"U{r}", f"={g}ROUND(P{r}*G{r}*Q{r}/1000*(I{r}+J{r})/12,2))", EUR2)
        cell(ca, f"V{r}", f"={g}ROUND(P{r}*G{r}*Q{r}/1000*S{r}/12,2))", EUR2)
        cell(ca, f"W{r}", f"={g}ROUND(P{r}*G{r}*R{r}/1000*(S{r}+T{r})/12,2))", EUR2)
        cell(ca, f"X{r}", f"=W{r}-V{r}", EUR2)
        cell(ca, f"Y{r}", f"=U{r}-W{r}", EUR2)
        cell(ca, f"Z{r}", f'={g}VLOOKUP(E{r},{cat_rng},3,0))')
        cell(ca, f"AA{r}", f'=IF($A{r}="","",VLOOKUP(E{r},{cat_rng},4,0))')
        cell(ca, f"AB{r}", f'=IF(Z{r}="D",ROUND(P{r}*G{r}*Parametri!$B$4/1000*(S{r}+T{r})/12,2),0)', EUR2)
        cell(ca, f"AC{r}", f"=W{r}-AB{r}", EUR2)
        cell(ca, f"AD{r}", f'=IF(Z{r}="D",ROUND(P{r}*G{r}*Parametri!$B$4/1000*S{r}/12,2),0)', EUR2)
        cell(ca, f"AE{r}", f"=AB{r}-AD{r}", EUR2)
        cell(ca, f"AF{r}", f'=IF($A{r}="","",VLOOKUP(D{r},{com_rng},2,0))')
    tr = last + 1
    cell(ca, f"A{tr}", "TOTALE", bold=True)
    for c_ in ["U", "V", "W", "X", "Y", "AB", "AC", "AD", "AE"]:
        cell(ca, f"{c_}{tr}", f"=SUM({c_}{first}:{c_}{last})", EUR2, bold=True)
    for c_ in range(1, 33):
        ca.cell(row=tr, column=c_).fill = FILL_TOTAL
    dv = DataValidation(type="list", formula1='"Sì,No"')
    ca.add_data_validation(dv)
    dv.add(f"H{first}:H{last}")
    ca.freeze_panes = "D5"

    R = lambda c_: f"'Calcolo IMU'!${c_}${first}:${c_}${last}"  # noqa: E731

    # Riepilogo F24
    f24 = wb.create_sheet("Riepilogo F24", 2)
    title(f24, f"Riepilogo versamenti F24 – Sezione IMU {ANNO}")
    widths(f24, [8, 12, 12, 12, 16, 16, 16])
    header(f24, 4, ["Società", "Comune", "Codice ente", "Codice tributo", "Acconto 16/06", "Saldo 16/12", "Totale"])
    combos = []
    for sc in [s[0] for s in SOCIETA]:
        for com, *_x in COMUNI:
            imms = [x for x in IMMOBILI if x[1] == sc and x[3] == com]
            if not imms:
                continue
            groups = {dict((c[0], c[2]) for c in CATEGORIE)[x[5]] for x in imms}
            if groups - {"D"}:
                combos.append((sc, com, "3918"))
            if "D" in groups:
                combos.append((sc, com, "3925"))
                combos.append((sc, com, "3930"))
    r = 5
    for sc, com, ct in combos:
        cell(f24, f"A{r}", sc)
        cell(f24, f"B{r}", com)
        cell(f24, f"C{r}", f"=VLOOKUP(B{r},{com_rng},2,0)")
        cell(f24, f"D{r}", ct)
        crit = f'{R("A")},A{r},{R("D")},B{r}'
        if ct == "3925":
            acc = f'=SUMIFS({R("AD")},{crit},{R("Z")},"D")'
            sal = f'=SUMIFS({R("AE")},{crit},{R("Z")},"D")'
        elif ct == "3930":
            acc = f'=SUMIFS({R("V")},{crit},{R("Z")},"D")-SUMIFS({R("AD")},{crit},{R("Z")},"D")'
            sal = f'=SUMIFS({R("X")},{crit},{R("Z")},"D")-SUMIFS({R("AE")},{crit},{R("Z")},"D")'
        else:
            acc = f'=SUMIFS({R("V")},{crit},{R("AA")},"3918")'
            sal = f'=SUMIFS({R("X")},{crit},{R("AA")},"3918")'
        cell(f24, f"E{r}", acc, EUR2)
        cell(f24, f"F{r}", sal, EUR2)
        cell(f24, f"G{r}", f"=E{r}+F{r}", EUR2)
        r += 1
    last_f = r - 1
    r += 1
    for sc, name, *_x in SOCIETA:
        cell(f24, f"A{r}", f"Totale {name}", bold=True)
        for c_ in "EFG":
            cell(f24, f"{c_}{r}", f'=SUMIFS({c_}$5:{c_}${last_f},$A$5:$A${last_f},"{sc}")', EUR2, bold=True)
        for c_ in range(1, 8):
            f24.cell(row=r, column=c_).fill = FILL_SECTION
        r += 1
    cell(f24, f"A{r}", "TOTALE GENERALE", bold=True)
    for c_ in "EFG":
        cell(f24, f"{c_}{r}", f"=SUM({c_}5:{c_}{last_f})", EUR2, bold=True)
    for c_ in range(1, 8):
        f24.cell(row=r, column=c_).fill = FILL_TOTAL
    f24.cell(row=r + 2, column=1, value="Controllo: il totale deve coincidere con 'Calcolo IMU' (IMU annua dovuta) →").font = F_SUB
    cell(f24, f"G{r + 2}", f"='Calcolo IMU'!W{tr}", EUR2)

    # Budget vs Dichiarazione
    bd = wb.create_sheet("Budget vs Dichiarazione", 3)
    title(bd, "IMU piena (budget) vs IMU dovuta (dichiarazione)",
          "Il budget di fine anno considera l'IMU piena; la dichiarazione la cristallizza sui mesi effettivi di locazione")
    widths(bd, [26, 14, 16, 16, 16, 12])
    header(bd, 4, ["Società", "Comune", "IMU piena (budget)", "IMU dovuta", "Risparmio", "Risparmio %"])
    r = 5
    for sc, name, *_x in SOCIETA:
        for com, *_y in COMUNI:
            if not any(x[1] == sc and x[3] == com for x in IMMOBILI):
                continue
            cell(bd, f"A{r}", sc)
            cell(bd, f"B{r}", com)
            crit = f'{R("A")},A{r},{R("D")},B{r}'
            cell(bd, f"C{r}", f"=SUMIFS({R('U')},{crit})", EUR2)
            cell(bd, f"D{r}", f"=SUMIFS({R('W')},{crit})", EUR2)
            cell(bd, f"E{r}", f"=C{r}-D{r}", EUR2)
            cell(bd, f"F{r}", f"=IFERROR(E{r}/C{r},0)", PCT1)
            r += 1
    last_b = r - 1
    r += 1
    for sc, name, *_x in SOCIETA:
        cell(bd, f"A{r}", name, bold=True)
        for c_ in "CDE":
            cell(bd, f"{c_}{r}", f'=SUMIFS({c_}$5:{c_}${last_b},$A$5:$A${last_b},"{sc}")', EUR2, bold=True)
        cell(bd, f"F{r}", f"=IFERROR(E{r}/C{r},0)", PCT1, bold=True)
        r += 1
    cell(bd, f"A{r}", "TOTALE", bold=True)
    for c_ in "CDE":
        cell(bd, f"{c_}{r}", f"=SUM({c_}5:{c_}{last_b})", EUR2, bold=True)
    cell(bd, f"F{r}", f"=IFERROR(E{r}/C{r},0)", PCT1, bold=True)
    for c_ in range(1, 7):
        bd.cell(row=r, column=c_).fill = FILL_TOTAL
    bch = BarChart()
    bch.title = "IMU piena vs dovuta per Comune"
    bch.height, bch.width = 8, 18
    bch.add_data(Reference(bd, min_col=3, max_col=4, min_row=4, max_row=last_b), titles_from_data=True)
    bch.set_categories(Reference(bd, min_col=2, max_col=2, min_row=5, max_row=last_b))
    bd.add_chart(bch, f"A{r + 3}")

    # Calcolatore rapido
    q = wb.create_sheet("Calcolatore rapido", 4)
    title(q, "Calcolatore IMU rapido – singolo immobile", "Inserisci i dati nelle celle gialle")
    widths(q, [36, 18, 4, 50])
    inputs = [
        ("Categoria catastale", "A/10", None),
        ("Rendita catastale", 2450, EUR2),
        ("Comune", "Napoli", None),
        ("Quota di possesso", 1, "0%"),
        ("Mesi imponibili (anno)", 12, "0"),
        ("Mesi imponibili 1° semestre", 6, "0"),
        ("Riduzione (50% inagibile / storico)", 0, "0%"),
    ]
    for i, (lab, v, fmt) in enumerate(inputs):
        r = 4 + i
        cell(q, f"A{r}", lab, bold=True)
        cell(q, f"B{r}", v, fmt, inp=True)
    outs = [
        ("Rendita rivalutata (+5%)", "=B5*(1+Parametri!$B$5)", EUR2),
        ("Moltiplicatore", f"=VLOOKUP(B4,{cat_rng},2,0)", "0"),
        ("Base imponibile", "=B12*B13*(1-B10)", EUR2),
        ("Aliquota anno precedente ‰", f"=VLOOKUP(B6,{com_rng},3,0)", "0.00"),
        ("Aliquota anno corrente ‰", f"=VLOOKUP(B6,{com_rng},4,0)", "0.00"),
        ("IMU annua dovuta", "=ROUND(B14*B7*B16/1000*B8/12,2)", EUR2),
        ("Acconto 16 giugno", "=ROUND(B14*B7*B15/1000*B9/12,2)", EUR2),
        ("Saldo 16 dicembre", "=B17-B18", EUR2),
        ("Gruppo catastale", f"=VLOOKUP(B4,{cat_rng},3,0)", None),
        ("di cui quota Stato (3925)", '=IF(B20="D",ROUND(B14*B7*Parametri!$B$4/1000*B8/12,2),0)', EUR2),
        ("di cui quota Comune", "=B17-B21", EUR2),
        ("Codice tributo Comune", f"=VLOOKUP(B4,{cat_rng},4,0)", None),
    ]
    for i, (lab, f, fmt) in enumerate(outs):
        r = 12 + i
        cell(q, f"A{r}", lab, bold=True)
        c = cell(q, f"B{r}", f, fmt, bold=lab.startswith("IMU"))
        if lab.startswith("IMU"):
            c.fill = FILL_TOTAL
    q["D4"] = "Categorie e Comuni disponibili: vedi foglio Parametri."
    q["D4"].font = F_SUB
    dvc = DataValidation(type="list", formula1=f"=Parametri!$D$8:$D${7 + len(CATEGORIE)}")
    dvm = DataValidation(type="list", formula1=f"=Parametri!$I$8:$I${7 + len(COMUNI)}")
    q.add_data_validation(dvc)
    q.add_data_validation(dvm)
    dvc.add("B4")
    dvm.add("B6")

    wb.save(path)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "."
    os.makedirs(out, exist_ok=True)
    build_cash_flow(os.path.join(out, f"Cash_Flow_Previsionale_{ANNO}.xlsx"))
    build_imu(os.path.join(out, f"Calcolatore_IMU_{ANNO}.xlsx"))
    print("ok")
